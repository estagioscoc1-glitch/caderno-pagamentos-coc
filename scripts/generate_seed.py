from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("/workspace/scratch/4157471acb05/upload/Relacao_de_Alunos_2026_2_Completa(4).xlsx")
TARGET = Path("/workspace/sites/caderno-pagamentos-coc/app/lib/seed-data.ts")

CLASS_NAMES = {
    "Enf 1 Noturno": "ENF 1 NOT",
    "Enf 1 Sábado": "ENF 1 SÁB",
    "Instrumentação 1": "INSTRUMENTAÇÃO",
    "Radiologia 1": "RAD 1 NOT",
    "Segurança 1": "TST 1 NOT",
    "Enf 1 Matutino": "ENF 1 MAT",
    "Enf 2 Matutino": "ENF 2 MAT",
    "Enf 3 EAD": "ENF 3 EAD",
    "Enf 2 Noturno": "ENF 2 NOT",
    "Enf 3 Noturno": "ENF 3 NOT",
    "Radiologia 2": "RAD 2 NOT",
    "Segurança 2": "TST 2 NOT",
}


def course_for(sheet: str) -> str:
    if sheet.startswith("Enf"):
        return "Técnico em Enfermagem"
    if sheet.startswith("Instrumentação"):
        return "Instrumentação Cirúrgica"
    if sheet.startswith("Radiologia"):
        return "Técnico em Radiologia"
    return "Técnico em Segurança do Trabalho"


def clean_name(name: str) -> str:
    name = re.sub(r"\s*\(?des\)?\s*$", "", name, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", name).strip()


workbook = load_workbook(SOURCE, data_only=True, read_only=True)
students = []

for sheet in workbook.worksheets:
    if sheet.title not in CLASS_NAMES:
        continue
    for row in sheet.iter_rows(min_row=3, values_only=True):
        values = list(row) + [None] * 7
        _, registration, name, module, shift, phone, source_status = values[:7]
        if not registration or not name:
            continue
        source = str(source_status or "").strip().upper()
        withdrawn = source in {"DES", "DESISTENTE", "CANCELOU", "CANCELADO"}
        september = "paid" if source == "PG" else "pending"
        students.append(
            {
                "id": f"student-{str(registration).strip()}",
                "registration": str(registration).strip(),
                "name": clean_name(str(name)),
                "phone": str(phone or "").strip(),
                "className": CLASS_NAMES[sheet.title],
                "course": course_for(sheet.title),
                "module": str(module or "").strip(),
                "shift": str(shift or "").strip(),
                "situation": "withdrawn" if withdrawn else "active",
                "situationChangedAt": "2026-09-01T12:00:00.000Z" if withdrawn else None,
                "payments": {
                    "2026-09": september,
                    "2026-10": "pending",
                    "2026-11": "pending",
                    "2026-12": "pending",
                    "2027-01": "pending",
                    "2027-02": "pending",
                    "insurance": "pending",
                    "initial-kit": "pending",
                    "enrollment": "pending",
                    "dependency": "pending",
                },
            }
        )

students.sort(key=lambda item: (item["className"], item["name"]))

header = '''export const PAYMENT_COLUMNS = [
  { id: "2026-09", label: "Setembro/26", shortLabel: "Set/26", kind: "month" },
  { id: "2026-10", label: "Outubro/26", shortLabel: "Out/26", kind: "month" },
  { id: "2026-11", label: "Novembro/26", shortLabel: "Nov/26", kind: "month" },
  { id: "2026-12", label: "Dezembro/26", shortLabel: "Dez/26", kind: "month" },
  { id: "2027-01", label: "Janeiro/27", shortLabel: "Jan/27", kind: "month" },
  { id: "2027-02", label: "Fevereiro/27", shortLabel: "Fev/27", kind: "month" },
  { id: "insurance", label: "Seguro", shortLabel: "Seguro", kind: "extra" },
  { id: "initial-kit", label: "Kit inicial", shortLabel: "Kit", kind: "extra" },
  { id: "enrollment", label: "Matrícula", shortLabel: "Matrícula", kind: "extra" },
  { id: "dependency", label: "Dependência", shortLabel: "Dep.", kind: "extra" },
] as const;

export type PaymentColumnId = (typeof PAYMENT_COLUMNS)[number]["id"];
export type PaymentStatus = "paid" | "pending";
export type StudentSituation = "active" | "withdrawn";

export type Student = {
  id: string;
  registration: string;
  name: string;
  phone: string;
  className: string;
  course: string;
  module: string;
  shift: string;
  situation: StudentSituation;
  situationChangedAt: string | null;
  payments: Record<PaymentColumnId, PaymentStatus>;
};

export type MessageLog = {
  id: string;
  studentId: string;
  columnId: PaymentColumnId;
  sentAt: string;
  text: string;
};

export type BookState = {
  version: 1;
  students: Student[];
  messageTemplate: string;
  messageHistory: MessageLog[];
  updatedAt: string;
};

export const CLASS_ORDER = [
  "ENF 1 NOT",
  "ENF 1 SÁB",
  "INSTRUMENTAÇÃO",
  "RAD 1 NOT",
  "TST 1 NOT",
  "ENF 1 MAT",
  "ENF 2 MAT",
  "ENF 3 EAD",
  "ENF 2 NOT",
  "ENF 3 NOT",
  "RAD 2 NOT",
  "TST 2 NOT",
] as const;

'''

state = {
    "version": 1,
    "students": students,
    "messageTemplate": "Olá, {nome}! Identificamos que {item} está pendente em nosso caderno de pagamentos. Por favor, entre em contato com a escola para regularização. Colégio Oswaldo Cruz.",
    "messageHistory": [],
    "updatedAt": "2026-09-10T12:00:00.000Z",
}
TARGET.parent.mkdir(parents=True, exist_ok=True)
TARGET.write_text(header + "export const INITIAL_STATE: BookState = " + json.dumps(state, ensure_ascii=False, indent=2) + ";\n", encoding="utf-8")
print(f"generated {TARGET} with {len(students)} students")
